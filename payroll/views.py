import io
import json
import math
import uuid
from datetime import date
from decimal import Decimal, InvalidOperation

from django.contrib import messages
from django.core.exceptions import PermissionDenied
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Sum
from django.http import Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from accounts.decorators import permission_required
from core.models import notify
from core.support import dec, money
from employees.models import Employee, Site

from . import rates
from leaveot.overtime import unpaid_watchlist

from .calculator import PayrollCalculator
from .models import ContributionRate, PayrollDeduction, PayrollItem, PayrollPeriod
from .runner import PayrollRunner


def _require_rates(request):
    if not request.user.can("manage payroll rates"):
        raise PermissionDenied


def _items_for(period):
    return sorted(period.items.select_related("employee", "adjusted_by"), key=lambda i: (i.employee.last_name.lower(), i.employee.first_name.lower()))


# ---- register ----

@permission_required("run payroll")
def index(request):
    runner = PayrollRunner()
    runner.roll_forward()
    periods = list(PayrollPeriod.objects.order_by("-period_start"))
    today = timezone.now().date()
    selected = None
    if request.GET.get("period", "").isdigit():
        selected = next((p for p in periods if p.id == int(request.GET["period"])), None)
    if selected is None:
        selected = next((p for p in periods if p.period_start <= today <= p.period_end), periods[0] if periods else None)
    items = _items_for(selected) if selected else []
    by_card = sum(1 for i in items if i.employee.pays_by_card)
    gross = float(sum(i.gross_pay for i in items)) or 1.0
    ded_total = float(sum(i.total_deductions for i in items)) or 1.0
    earn = [("Basic", sum(i.basic_pay for i in items)), ("Overtime", sum(i.overtime_pay for i in items)), ("Allowances", sum(i.allowances for i in items)), ("Holiday / night", sum(i.holiday_pay + i.night_diff_pay for i in items))]
    ded = [("SSS", sum(i.sss_deduction for i in items)), ("PhilHealth", sum(i.philhealth_deduction for i in items)), ("Pag-IBIG", sum(i.pagibig_deduction for i in items)), ("Tax", sum(i.withholding_tax for i in items)),
           ("Absences / half-day", sum(i.absences_deduction + i.half_day_deduction for i in items)), ("Late / other", sum(i.late_undertime_deduction + i.other_deductions for i in items)),
           ("Loans", sum(i.loan_deduction for i in items)), ("Missing items", sum(i.missing_item_deduction for i in items))]
    earn_colors = ["bg-emerald-500", "bg-emerald-400", "bg-emerald-300", "bg-emerald-200"]
    ded_colors = ["bg-accent-600", "bg-accent-400", "bg-amber-400", "bg-rose-500", "bg-rose-300", "bg-slate-400", "bg-sky-500", "bg-violet-500"]
    pay_day = runner.status()
    ot_watch = unpaid_watchlist(selected) if (selected and not selected.is_released and not selected.is_closed) else None
    return render(request, "payroll/index.html", {
        "periods": periods, "selected": selected, "items": items, "pay_day": pay_day, "ot_watch": ot_watch,
        "can_edit": request.user.can("manage payroll rates"), "by_card": by_card, "cash_count": len(items) - by_card,
        "totals": {"gross": sum(i.gross_pay for i in items), "deductions": sum(i.total_deductions for i in items), "net": sum(i.net_pay for i in items)},
        "kept_pct": round(100 - float(sum(i.total_deductions for i in items)) / gross * 100, 1), "ded_pct": round(float(sum(i.total_deductions for i in items)) / gross * 100, 1),
        "earn": [{"label": l, "value": v, "pct": round(float(v) / gross * 100, 1), "color": earn_colors[k]} for k, (l, v) in enumerate(earn)],
        "ded": [{"label": l, "value": v, "pct": round(float(v) / gross * 100, 1), "share": int(round(float(v) / ded_total * 100)), "color": ded_colors[k]} for k, (l, v) in enumerate(ded)],
        "search_rows": ([{"id": i.id, "name": i.employee.full_name, "no": i.employee.employee_no or "", "method": "Card" if i.employee.pays_by_card else "Cash", "url": reverse("payroll.print", args=[selected.id]) + f"?employee={i.employee_id}"} for i in items] if selected else []),
        "back_url": reverse("dashboard"), "back_label": "Back to Dashboard",
    })


@permission_required("run payroll")
@require_POST
def generate(request, pk):
    period = get_object_or_404(PayrollPeriod, pk=pk)
    if period.is_closed:
        messages.error(request, "This period is closed.")
        return redirect(f"{reverse('payroll.index')}?period={period.id}")
    force = request.POST.get("force") in ("1", "true", "on")
    n = PayrollRunner().generate(period, str(request.user.id), force)
    kept = period.items.filter(adjusted_at__isnull=False).count()
    messages.success(request, f"Payroll computed for {n} employees." + (f" {kept} manually adjusted line(s) were kept as-is." if kept and not force else ""))
    return redirect(f"{reverse('payroll.index')}?period={period.id}")


@permission_required("run payroll")
@require_POST
def create_period(request):
    latest = PayrollPeriod.objects.order_by("-period_end").first()
    period = latest.next() if latest else PayrollPeriod.ensure_for(timezone.now().date())
    messages.success(request, f"Period {period.label()} is ready.")
    return redirect(f"{reverse('payroll.index')}?period={period.id}")


@permission_required("run payroll")
@require_POST
def release(request, pk):
    _require_rates(request)
    period = get_object_or_404(PayrollPeriod, pk=pk)
    if not period.items.exists():
        messages.error(request, "Generate the payroll before releasing it.")
        return redirect(f"{reverse('payroll.index')}?period={period.id}")
    if period.is_released:
        messages.error(request, "This payroll was already released.")
        return redirect(f"{reverse('payroll.index')}?period={period.id}")
    now = timezone.now()
    period.status, period.closed_at, period.closed_by = "closed", period.closed_at or now, period.closed_by or request.user
    period.released_at, period.released_by = now, request.user
    period.save()
    for item in period.items.select_related("employee__user"):
        u = item.employee.user
        if u:
            how = "credited to your card" if item.employee.pays_by_card else "released in cash with your printed payslip"
            notify(u, kind="payroll", title="Your salary has been released", message=f"{period.label()} — {how}. Net {money(item.net_pay)}.", url=reverse("payroll.show", args=[item.id]))
    messages.success(request, f"Payroll released ({period.release_timing()}). Every employee can now see and download their payslip.")
    return redirect(f"{reverse('payroll.index')}?period={period.id}")


@permission_required("run payroll")
@require_POST
def close(request, pk):
    _require_rates(request)
    period = get_object_or_404(PayrollPeriod, pk=pk)
    if not period.items.exists():
        messages.error(request, "Generate the payroll before closing the period.")
        return redirect(f"{reverse('payroll.index')}?period={period.id}")
    period.status, period.closed_at, period.closed_by = "closed", timezone.now(), request.user
    period.save()
    messages.success(request, f"Period {period.label()} closed. Payslips are final.")
    return redirect(f"{reverse('payroll.index')}?period={period.id}")


@permission_required("run payroll")
def export(request, pk):
    period = get_object_or_404(PayrollPeriod, pk=pk)
    items = _items_for(period)
    headers = ["Employee No", "Employee", "Payout", "Card / account no.", "Basic pay", "Overtime", "Night diff", "Holiday", "Allowances", "Gross",
               "Late / undertime", "Absences", "Half-day", "SSS", "PhilHealth", "Pag-IBIG", "Withholding tax", "Other deductions", "Loan", "Missing item", "Total deductions", "Net pay", "Adjusted", "Remarks"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Payroll"
    title = f"Brite TSI — Payroll register · {period.label()} · pay date {period.pay_date:%b %-d, %Y}" if period.pay_date else f"Brite TSI — Payroll register · {period.label()}"
    if period.is_released:
        title += f" · released {period.released_at:%b %-d, %Y}"
    ws["A1"] = title
    ws.merge_cells("A1:X1")
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    ws.append(headers)
    for c in range(1, 25):
        cell = ws.cell(row=3, column=c)
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", start_color="FF0E7490")
        ws.column_dimensions[get_column_letter(c)].width = 16
    for i in items:
        e = i.employee
        ws.append([e.employee_no, e.full_name, "Card" if e.pays_by_card else "Cash", e.bank_account_no,
                   float(i.basic_pay), float(i.overtime_pay), float(i.night_diff_pay), float(i.holiday_pay), float(i.allowances), float(i.gross_pay),
                   float(i.late_undertime_deduction), float(i.absences_deduction), float(i.half_day_deduction), float(i.sss_deduction), float(i.philhealth_deduction),
                   float(i.pagibig_deduction), float(i.withholding_tax), float(i.other_deductions), float(i.loan_deduction), float(i.missing_item_deduction), float(i.total_deductions), float(i.net_pay),
                   "Yes" if i.is_adjusted else "", i.remarks])
    last = 3 + len(items)
    t = last + 1
    ws.cell(row=t, column=2, value="TOTAL").font = Font(bold=True)
    for c in range(5, 23):
        col = get_column_letter(c)
        ws.cell(row=t, column=c, value=f"=SUM({col}4:{col}{last})").font = Font(bold=True)
    for r in range(4, t + 1):
        for c in range(5, 23):
            ws.cell(row=r, column=c).number_format = "#,##0.00"
    ws.freeze_panes = "C4"
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="payroll-{period.period_start}-to-{period.period_end}.xlsx"'
    return resp


@permission_required("run payroll")
def print_batch(request, pk):
    period = get_object_or_404(PayrollPeriod, pk=pk)
    method = request.GET.get("method") if request.GET.get("method") in ("card", "cash") else None
    employee_id = int(request.GET["employee"]) if request.GET.get("employee", "").isdigit() else None
    items = _items_for(period)
    if employee_id:
        items = [i for i in items if i.employee_id == employee_id]
        if not items:
            raise Http404("No payslip for that employee in this period.")
    if method:
        items = [i for i in items if i.employee.payout_method == method]
    return render(request, "payroll/print.html", {"period": period, "items": items, "method": method, "single": items[0] if employee_id else None})


# ---- payslips ----

@permission_required("view own payslip")
def mine(request):
    employee = getattr(request.user, "employee", None)
    items = []
    if employee:
        items = [i for i in employee.payroll_items.select_related("payroll_period") if i.payroll_period.is_released]
        items.sort(key=lambda i: i.payroll_period.period_start, reverse=True)
    year = timezone.now().date().year
    ytd = [i for i in items if i.payroll_period.period_start.year == year]
    summary = {"last": items[0] if items else None, "ytd_net": sum((i.net_pay for i in ytd), Decimal(0)), "ytd_count": len(ytd),
               "method": "Card" if (employee and employee.pays_by_card) else "Cash", "next": PayrollPeriod.cutoff_for(timezone.now().date())["end"]}
    return render(request, "payroll/mine.html", {"items": items, "summary": summary, "year": year})


def show(request, pk):
    if not request.user.is_authenticated:
        raise PermissionDenied
    item = get_object_or_404(PayrollItem.objects.select_related("employee", "payroll_period"), pk=pk)
    employee = getattr(request.user, "employee", None)
    is_owner = bool(employee) and item.employee_id == employee.id
    is_staff = request.user.can_any("view all payslips", "run payroll")
    if not (is_owner or is_staff):
        raise PermissionDenied
    if not is_staff and not item.payroll_period.is_released:
        raise Http404("This payslip is not released yet.")
    return render(request, "payroll/show.html", {"item": item, "period": item.payroll_period, "back_url": reverse("payroll.index") if is_staff else reverse("payroll.mine"), "back_label": "Back"})


@permission_required("run payroll")
def salary_history(request, pk):
    employee = get_object_or_404(Employee, pk=pk)
    items = sorted(employee.payroll_items.select_related("payroll_period"), key=lambda i: i.payroll_period.period_start, reverse=True)
    return render(request, "payroll/salary_history.html", {"employee": employee, "items": items, "total_net": sum((i.net_pay for i in items), Decimal(0)), "total_gross": sum((i.gross_pay for i in items), Decimal(0)), "total_deductions": sum((i.total_deductions for i in items), Decimal(0)), "back_url": reverse("employees.index"), "back_label": "Back to Employees"})


# ---- lines ----

@permission_required("run payroll")
def edit_line(request, pk):
    _require_rates(request)
    item = get_object_or_404(PayrollItem.objects.select_related("employee", "payroll_period", "adjusted_by"), pk=pk)
    payments = list(item.deduction_payments.select_related("deduction__site"))
    earn = [("basic_pay", "Basic pay"), ("overtime_pay", "Overtime pay"), ("night_diff_pay", "Night differential"), ("holiday_pay", "Holiday pay"), ("allowances", "Allowances")]
    ded = [("late_undertime_deduction", "Late / undertime"), ("absences_deduction", "Absences"), ("half_day_deduction", "Half-day leave"), ("sss_deduction", "SSS"), ("philhealth_deduction", "PhilHealth"), ("pagibig_deduction", "Pag-IBIG"), ("withholding_tax", "Withholding tax"), ("other_deductions", "Other deductions")]
    vals = {k: float(getattr(item, k)) for k, _ in earn + ded}
    return render(request, "payroll/edit.html", {"item": item, "period": item.payroll_period, "payments": payments, "earn": earn, "ded": ded, "vals": vals,
                                                  "fixed_deductions": float(item.loan_deduction) + float(item.missing_item_deduction),
                                                  "errors": request.session.pop("form_errors", None) or {}, "back_url": f"{reverse('payroll.index')}?period={item.payroll_period_id}", "back_label": "Back to Payroll"})


@permission_required("run payroll")
@require_POST
def update_line(request, pk):
    _require_rates(request)
    item = get_object_or_404(PayrollItem.objects.select_related("employee", "payroll_period"), pk=pk)
    period = item.payroll_period
    if period.is_closed:
        messages.error(request, "This period is closed.")
        return redirect(f"{reverse('payroll.index')}?period={period.id}")
    errors = {}
    for f in PayrollItem.EDITABLE:
        try:
            v = Decimal(request.POST.get(f, "0") or "0")
            if v < 0 or v > 9999999:
                raise InvalidOperation
            setattr(item, f, dec(v))
        except InvalidOperation:
            errors[f] = "Enter a valid amount."
    if errors:
        request.session["form_errors"] = errors
        return redirect("payroll.lines.edit", pk=pk)
    item.remarks = (request.POST.get("remarks") or "").strip()[:500] or None
    item.adjusted_at, item.adjusted_by = timezone.now(), request.user
    item.recompute_totals().save()
    messages.success(request, f"Payroll line for {item.employee.full_name} updated — net {money(item.net_pay)}.")
    return redirect(f"{reverse('payroll.index')}?period={period.id}")


@permission_required("run payroll")
@require_POST
def reset_line(request, pk):
    _require_rates(request)
    item = get_object_or_404(PayrollItem.objects.select_related("employee", "payroll_period"), pk=pk)
    period = item.payroll_period
    if period.is_closed:
        messages.error(request, "This period is closed.")
        return redirect(f"{reverse('payroll.index')}?period={period.id}")
    PayrollCalculator().calculate(item.employee, period, force=True)
    messages.success(request, "Line recomputed from attendance and rates.")
    return redirect(f"{reverse('payroll.index')}?period={period.id}")


# ---- rates ----

@permission_required("manage payroll rates")
def rates_index(request):
    active = Employee.objects.active()
    salaries = [float(s) for s in active.values_list("monthly_salary", flat=True)]
    sample = round(sum(salaries) / len(salaries), 2) if salaries else 20000.0
    brackets = list(ContributionRate.objects.order_by("contribution_type", "min_salary"))
    by_type = {}
    for b in brackets:
        by_type.setdefault(b.contribution_type, []).append(b)
    return render(request, "payroll/rates.html", {
        "rates": rates.all_rates(), "labels": rates.LABELS, "brackets": brackets, "by_type": by_type, "type_labels": {"sss": "SSS", "philhealth": "PhilHealth", "pagibig": "Pag-IBIG"},
        "sample": sample, "employees": active.count(), "payroll_monthly": float(active.aggregate(s=Sum("monthly_salary"))["s"] or 0),
        "rates_js": rates.all_rates(),
        "earn_rows": [("basic_cutoff_percent", "%", "of the monthly salary paid each cutoff (50 = semi-monthly)"), ("allowance_percent", "%", "of basic pay added to everyone as allowance (0 = none)"), ("ot_regular_percent", "%", "of the hourly rate per approved OT hour"), ("ot_rest_day_percent", "%", "of the hourly rate on rest days"), ("ot_holiday_percent", "%", "of the hourly rate on regular holidays")],
        "bracket_js": ([{"id": b.id, "type": b.contribution_type, "min": float(b.min_salary), "max": None if b.max_salary is None else float(b.max_salary), "emp": round(float(b.employee_rate) * 100, 2), "er": round(float(b.employer_rate) * 100, 2)} for b in brackets]),
        "back_url": reverse("payroll.index"), "back_label": "Back to Payroll",
    })


@permission_required("manage payroll rates")
@require_POST
def rates_update(request):
    limits = {"basic_cutoff_percent": (1, 100), "working_days_per_month": (1, 31), "hours_per_day": (1, 24), "allowance_percent": (0, 100), "ot_regular_percent": (100, 400),
              "ot_rest_day_percent": (100, 400), "ot_holiday_percent": (100, 400), "withholding_tax_percent": (0, 50), "pagibig_monthly_cap": (0, 100000)}
    values = {}
    for k, (lo, hi) in limits.items():
        try:
            v = float(request.POST.get(f"rates[{k}]", ""))
        except ValueError:
            messages.error(request, f"{rates.LABELS[k]} must be a number.")
            return redirect("payroll.rates")
        if not lo <= v <= hi:
            messages.error(request, f"{rates.LABELS[k]} must be between {lo} and {hi}.")
            return redirect("payroll.rates")
        values[k] = int(v) if k == "working_days_per_month" else v
    with transaction.atomic():
        rates.set_rates(values)
        for b in ContributionRate.objects.all():
            p = f"brackets[{b.id}]"
            if f"{p}[employee_rate]" not in request.POST:
                continue
            try:
                b.employee_rate = round(float(request.POST.get(f"{p}[employee_rate]", 0)) / 100, 4)
                b.employer_rate = round(float(request.POST.get(f"{p}[employer_rate]", 0)) / 100, 4)
                b.min_salary = Decimal(request.POST.get(f"{p}[min_salary]", "0") or "0")
                mx = request.POST.get(f"{p}[max_salary]", "")
                b.max_salary = Decimal(mx) if mx not in ("", None) else None
                b.save()
            except (ValueError, InvalidOperation):
                messages.error(request, "A contribution bracket has an invalid number.")
                return redirect("payroll.rates")
    messages.success(request, "Payroll rates saved. They apply the next time a payroll is computed or recalculated (manually adjusted lines are kept).")
    return redirect("payroll.rates")


# ---- loans & missing items ----

def _refresh_computed_line(deduction):
    """Recompute unreleased, already-computed periods the deduction lands in — unless the line was edited by hand."""
    calc = PayrollCalculator()
    for period in PayrollPeriod.objects.filter(released_at__isnull=True, generated_at__isnull=False, period_end__gte=deduction.starts_on).order_by("period_start"):
        item = PayrollItem.objects.filter(employee_id=deduction.employee_id, payroll_period=period).first()
        if item and not item.is_adjusted:
            calc.calculate(deduction.employee, period)


@permission_required("run payroll")
def deductions_index(request):
    kind = request.GET.get("type") if request.GET.get("type") in PayrollDeduction.TYPES else None
    status = request.GET.get("status") if request.GET.get("status") in ("active", "paid", "cancelled") else "active"
    qs = PayrollDeduction.objects.select_related("employee", "site", "created_by").filter(status=status).order_by("-id")
    if kind:
        qs = qs.filter(type=kind)
    page = Paginator(qs, 25).get_page(request.GET.get("page"))
    active = PayrollDeduction.objects.filter(status=PayrollDeduction.ACTIVE)
    stats = {
        "loans": active.filter(type=PayrollDeduction.LOAN).count(), "loan_balance": active.filter(type=PayrollDeduction.LOAN).aggregate(s=Sum("balance"))["s"] or 0,
        "missing": active.filter(type=PayrollDeduction.MISSING_ITEM).count(), "missing_balance": active.filter(type=PayrollDeduction.MISSING_ITEM).aggregate(s=Sum("balance"))["s"] or 0,
        "next_cutoff": sum((d.next_installment() for d in active), Decimal(0)),
    }
    stats["outstanding"] = Decimal(stats["loan_balance"]) + Decimal(stats["missing_balance"])
    old = request.session.pop("form_old", None) or {}
    return render(request, "payroll/deductions.html", {
        "deductions": page, "type": kind, "status": status, "stats": stats,
        "employees": Employee.objects.active().order_by("last_name", "first_name"), "sites": sorted(Site.objects.all(), key=lambda s: (0 if s.type == "office" else 1, s.name)),
        "next_cutoff": PayrollPeriod.cutoff_for(timezone.now().date())["end"], "can_manage": request.user.can("manage payroll rates"),
        "errors": request.session.pop("form_errors", None), "old": old, "old_people": old.get("employees") or [], "today": timezone.now().date().isoformat(),
        "back_url": reverse("payroll.index"), "back_label": "Back to Payroll",
    })


@permission_required("run payroll")
@require_POST
def deductions_store(request):
    _require_rates(request)
    P = request.POST
    kind = P.get("type")
    errors = {}
    if kind not in PayrollDeduction.TYPES:
        errors["type"] = "Pick loan or missing item."
    description = (P.get("description") or "").strip()[:255]
    if not description:
        errors["description"] = "Describe it."
    try:
        total = dec(Decimal(P.get("total_amount", "0") or "0"))
        if total <= 0 or total > 99999999:
            raise InvalidOperation
    except InvalidOperation:
        errors["total_amount"] = "Enter an amount above zero."
        total = Decimal(0)
    try:
        cutoffs = int(P.get("cutoffs", "1"))
        if not 1 <= cutoffs <= 120:
            raise ValueError
    except ValueError:
        errors["cutoffs"] = "Between 1 and 120 cutoffs."
        cutoffs = 1
    try:
        starts_on = date.fromisoformat(P.get("starts_on", ""))
    except ValueError:
        errors["starts_on"] = "Pick a start date."
        starts_on = None
    remarks = (P.get("remarks") or "").strip()[:1000] or None
    employee = site = None
    ids = []
    if kind == PayrollDeduction.LOAN:
        employee = Employee.objects.filter(pk=P.get("employee_id") or 0).first()
        if not employee:
            errors["employee_id"] = "Pick the employee who took the loan."
    elif kind == PayrollDeduction.MISSING_ITEM:
        site = Site.objects.filter(pk=P.get("site_id") or 0).first()
        if not site:
            errors["site_id"] = "Pick the site where the item went missing."
        ids = [int(i) for i in P.getlist("employees") if str(i).isdigit()]
        ids = list(dict.fromkeys(ids))
        if not ids or Employee.objects.filter(pk__in=ids).count() != len(ids):
            errors["employees"] = "Pick at least one employee responsible for the missing item."
    if errors:
        request.session["form_errors"], request.session["form_old"] = errors, {k: P.get(k) for k in ("type", "description", "total_amount", "cutoffs", "starts_on", "remarks", "employee_id", "site_id")} | {"employees": ids}
        return redirect("payroll.deductions")
    created = []
    with transaction.atomic():
        if kind == PayrollDeduction.LOAN:
            rows = [(employee.id, total)]
            group = None
        else:
            share = (total / len(ids)).quantize(Decimal("0.01"), rounding="ROUND_FLOOR")
            remainder = dec(total - share * len(ids))
            rows = [(eid, dec(share + (remainder if k == 0 else 0))) for k, eid in enumerate(ids)]
            group = str(uuid.uuid4())
        for eid, amount in rows:
            created.append(PayrollDeduction.objects.create(
                employee_id=eid, type=kind, site=site if kind == PayrollDeduction.MISSING_ITEM else None, group_id=group, description=description,
                incident_date=None, total_amount=amount, installment_amount=dec(Decimal(math.ceil(float(amount) / cutoffs * 100)) / 100), balance=amount,
                starts_on=starts_on, status=PayrollDeduction.ACTIVE, remarks=remarks, created_by=request.user,
            ))
    for d in created:
        _refresh_computed_line(d)
    if kind == PayrollDeduction.LOAN:
        messages.success(request, f"Loan recorded: {money(total)} over {cutoffs} cutoff(s).")
    else:
        messages.success(request, f"Missing item charged to {len(created)} employee(s): {money(total)} over {cutoffs} cutoff(s).")
    return redirect(f"{reverse('payroll.deductions')}?type={kind}")


@permission_required("run payroll")
@require_POST
def deductions_update(request, pk):
    _require_rates(request)
    d = get_object_or_404(PayrollDeduction.objects.select_related("employee"), pk=pk)
    if not d.is_active:
        messages.error(request, "Only an active balance can be changed.")
        return redirect("payroll.deductions")
    try:
        inst = dec(Decimal(request.POST.get("installment_amount", "0")))
        if inst <= 0:
            raise InvalidOperation
        starts = date.fromisoformat(request.POST.get("starts_on", ""))
    except (InvalidOperation, ValueError):
        messages.error(request, "Enter a valid installment and start date.")
        return redirect("payroll.deductions")
    d.installment_amount, d.starts_on, d.remarks = inst, starts, (request.POST.get("remarks") or "").strip()[:1000] or None
    d.save()
    _refresh_computed_line(d)
    messages.success(request, f"Updated. Next cutoff takes {money(d.next_installment())}.")
    return redirect(request.META.get("HTTP_REFERER") or "payroll.deductions")


@permission_required("run payroll")
@require_POST
def deductions_cancel(request, pk):
    _require_rates(request)
    d = get_object_or_404(PayrollDeduction.objects.select_related("employee"), pk=pk)
    if not d.is_active:
        messages.error(request, "This balance is not active.")
        return redirect("payroll.deductions")
    reason = (request.POST.get("reason") or "").strip()
    d.status, d.cancelled_at, d.cancelled_by = PayrollDeduction.CANCELLED, timezone.now(), request.user
    d.remarks = ((d.remarks + "\n") if d.remarks else "") + f"Cancelled: {reason}"
    d.save()
    _refresh_computed_line(d)
    messages.success(request, f"Cancelled — {money(d.balance)} will not be collected.")
    return redirect(request.META.get("HTTP_REFERER") or "payroll.deductions")
