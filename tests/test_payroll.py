import io
from datetime import date, datetime, time, timedelta
from decimal import Decimal

from django.core.management import call_command
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from attendance.models import AttendanceLog
from core.models import Notification
from employees.models import Employee
from leaveot.models import OvertimeRequest
from payroll import rates
from payroll.models import PayrollDeduction, PayrollItem, PayrollPeriod
from payroll.runner import PayrollRunner

from .base import HQ, AppTestCase


def punch_day(employee, d: date, time_in="08:30", time_out="17:30"):
    for kind, t in (("time_in", time_in), ("time_out", time_out)):
        AttendanceLog.objects.create(employee=employee, log_type=kind, logged_at=datetime.combine(d, datetime.strptime(t, "%H:%M").time()), latitude=HQ[0], longitude=HQ[1], location_status="verified_location", within_geofence=True, photo_path="x.jpg")


class PeriodTest(AppTestCase):
    def test_cutoffs_are_fixed_semi_monthly(self):
        p = PayrollPeriod.ensure_for(date(2026, 9, 3))
        self.assertEqual((p.period_start, p.period_end), (date(2026, 9, 1), date(2026, 9, 15)))
        q = p.next()
        self.assertEqual((q.period_start, q.period_end), (date(2026, 9, 16), date(2026, 9, 30)))
        self.assertEqual(q.next().period_start, date(2026, 10, 1))
        self.assertEqual(PayrollPeriod.ensure_for(date(2026, 2, 20)).period_end, date(2026, 2, 28))

    def test_ensure_for_is_idempotent(self):
        PayrollPeriod.ensure_for(date(2026, 9, 3))
        PayrollPeriod.ensure_for(date(2026, 9, 10))
        self.assertEqual(PayrollPeriod.objects.count(), 1)


class RunnerTest(AppTestCase):
    def test_roll_forward_fills_the_gap_up_to_the_next_cutoff(self):
        PayrollPeriod.ensure_for(date(2026, 7, 1))
        PayrollRunner().roll_forward(datetime(2026, 9, 20, 7, 0))
        ends = list(PayrollPeriod.objects.order_by("period_end").values_list("period_end", flat=True))
        self.assertEqual(ends[0], date(2026, 7, 15))
        self.assertIn(date(2026, 9, 30), ends)
        self.assertIn(date(2026, 8, 31), ends)

    def test_pay_day_reminder_goes_out_once_on_the_morning_of_pay_day(self):
        runner = PayrollRunner()
        pay_day = datetime(2026, 9, 15, 7, 0)
        result = runner.tick(pay_day)
        self.assertEqual(len(result["reminded"]), 1)
        self.assertTrue(Notification.objects.filter(user=self.admin, kind="payroll").exists())
        self.assertTrue(Notification.objects.filter(user=self.hr, kind="payroll").exists())
        self.assertFalse(Notification.objects.filter(user=self.tech, kind="payroll").exists())
        again = runner.tick(pay_day + timedelta(hours=2))
        self.assertEqual(again["reminded"], [])

    def test_status_tells_the_admin_what_payroll_needs_right_now(self):
        runner = PayrollRunner()
        self.assertEqual(runner.status(datetime(2026, 9, 3))["state"], "upcoming")
        runner.roll_forward(datetime(2026, 9, 15))
        self.assertEqual(runner.status(datetime(2026, 9, 15))["state"], "due")
        self.assertEqual(runner.status(datetime(2026, 9, 18))["state"], "overdue")
        period = PayrollPeriod.objects.get(period_end=date(2026, 9, 15))
        runner.generate(period, by="test")
        self.assertEqual(runner.status(datetime(2026, 9, 18))["state"], "computed")
        period.released_at, period.status = timezone.now(), "released"
        period.save()
        s = runner.status(datetime(2026, 9, 18))
        self.assertEqual((s["state"], s["pay_date"]), ("upcoming", date(2026, 9, 30)))

    def test_management_command(self):
        out = io.StringIO()
        call_command("payroll_tick", stdout=out)
        self.assertTrue(PayrollPeriod.objects.exists())


class CalculatorTest(AppTestCase):
    def setUp(self):
        self.period = PayrollPeriod.ensure_for(date(2026, 9, 1))  # Sep 1–15, 2026 (11 working days)

    def run_payroll(self):
        self.login(self.admin)
        return self.client.post(reverse("payroll.generate", args=[self.period.id]))

    def test_hr_can_run_but_not_release_or_edit(self):
        self.login(self.hr)
        self.assertEqual(self.client.post(reverse("payroll.generate", args=[self.period.id])).status_code, 302)
        self.assertEqual(self.client.post(reverse("payroll.release", args=[self.period.id])).status_code, 403)
        line = self.period.items.get(employee=self.employee)
        self.assertEqual(self.client.get(reverse("payroll.lines.edit", args=[line.id])).status_code, 403)

    def test_generate_creates_a_line_per_active_employee(self):
        r = self.run_payroll()
        self.assertRedirects(r, f"{reverse('payroll.index')}?period={self.period.id}", fetch_redirect_response=False)
        self.period.refresh_from_db()
        self.assertEqual(self.period.status, "processing")
        self.assertEqual(self.period.items.count(), Employee.objects.active().count())
        line = self.period.items.get(employee=self.employee)
        self.assertEqual(float(line.basic_pay), 12500.0)   # 25 000 × 50 %
        self.assertEqual(float(line.absences_deduction), 0.0)  # no attendance captured → nothing zeroed
        self.assertGreater(float(line.sss_deduction), 0)
        self.assertEqual(float(line.net_pay), float(line.gross_pay - line.total_deductions))

    def test_absences_are_deducted_once_attendance_exists(self):
        for d in (date(2026, 9, 1), date(2026, 9, 2), date(2026, 9, 3)):
            punch_day(self.employee, d)
        self.run_payroll()
        line = self.period.items.get(employee=self.employee)
        daily = float(self.employee.daily_rate)
        self.assertEqual(float(line.absences_deduction), round((11 - 3) * daily, 2))

    def test_late_arrival_is_deducted_per_minute_after_grace(self):
        punch_day(self.employee, date(2026, 9, 1), time_in="09:00")  # 30 min late, 15 min grace
        self.run_payroll()
        line = self.period.items.get(employee=self.employee)
        self.assertGreater(float(line.late_undertime_deduction), 0)

    def test_hand_edited_line_survives_a_rerun_unless_forced(self):
        self.run_payroll()
        line = self.period.items.get(employee=self.employee)
        data = {f: str(getattr(line, f)) for f in PayrollItem.EDITABLE}
        data["allowances"] = "500"
        data["remarks"] = "Transport allowance"
        r = self.client.post(reverse("payroll.lines.update", args=[line.id]), data)
        self.assertEqual(r.status_code, 302)
        line.refresh_from_db()
        self.assertEqual(float(line.allowances), 500.0)
        self.assertTrue(line.is_adjusted)
        self.client.post(reverse("payroll.generate", args=[self.period.id]))
        line.refresh_from_db()
        self.assertEqual(float(line.allowances), 500.0)
        self.client.post(reverse("payroll.lines.reset", args=[line.id]))
        line.refresh_from_db()
        self.assertEqual(float(line.allowances), 500.0)  # hand-entered allowances survive; attendance figures recompute
        self.assertFalse(line.is_adjusted)

    def test_release_issues_payslips_and_notifies_employees(self):
        self.run_payroll()
        r = self.client.post(reverse("payroll.release", args=[self.period.id]))
        self.assertEqual(r.status_code, 302)
        self.period.refresh_from_db()
        self.assertIsNotNone(self.period.released_at)
        self.assertTrue(self.period.is_closed)  # releasing closes the period: the figures are final
        self.assertTrue(Notification.objects.filter(user=self.tech, kind="payroll").exists())
        self.login(self.tech)
        r = self.client.get(reverse("payroll.mine"))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Sep 1")

    def test_close_locks_editing(self):
        self.run_payroll()
        self.client.post(reverse("payroll.release", args=[self.period.id]))
        self.client.post(reverse("payroll.close", args=[self.period.id]))
        self.period.refresh_from_db()
        self.assertTrue(self.period.is_closed)
        line = self.period.items.get(employee=self.employee)
        self.client.post(reverse("payroll.lines.update", args=[line.id]), {f: "1" for f in PayrollItem.EDITABLE})
        line.refresh_from_db()
        self.assertNotEqual(float(line.basic_pay), 1.0)

    def test_export_and_print(self):
        self.run_payroll()
        r = self.client.get(reverse("payroll.export", args=[self.period.id]))
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(io.BytesIO(r.content))
        self.assertGreaterEqual(wb.active.max_row, 2)
        r = self.client.get(reverse("payroll.print", args=[self.period.id]))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Payslip")
        item = self.period.items.get(employee=self.employee)
        self.assertEqual(self.client.get(reverse("payroll.show", args=[item.id])).status_code, 200)

    def test_employee_cannot_run_payroll(self):
        self.login(self.tech)
        self.assertEqual(self.client.post(reverse("payroll.generate", args=[self.period.id])).status_code, 403)
        self.assertEqual(self.client.get(reverse("payroll.index")).status_code, 403)



class OvertimePayTest(AppTestCase):
    """Approval is permission and a cap; what gets paid comes from the punches (and HR's 13h+ verdict)."""

    OT_DAY = date(2026, 9, 8)  # a Tuesday in the Sep 1–15 window paid by the Sep 16–30 cutoff

    def setUp(self):
        self.period = PayrollPeriod.ensure_for(date(2026, 9, 16))
        self.assertEqual(self.period.overtime_window(), (date(2026, 9, 1), date(2026, 9, 15)))
        self.ot = OvertimeRequest.objects.create(
            employee=self.employee, ot_date=self.OT_DAY, planned_start=time(17, 30), planned_end=time(20, 30), requested_hours=3,
            ot_type="regular", reason="Finish the pull-out at Site A", status="approved", approved_at=timezone.now(),
        )
        self.hourly = float(self.employee.daily_rate) / 8

    def ot_pay(self):
        self.login(self.admin)
        self.client.post(reverse("payroll.generate", args=[self.period.id]))
        self.ot.refresh_from_db()
        return float(self.period.items.get(employee=self.employee).overtime_pay)

    def verify(self, decision, remarks="Checked with the site lead"):
        out = self.employee.attendance_logs.get(log_type="time_out", logged_at__date=self.OT_DAY)
        self.login(self.hr)
        self.client.post(reverse("attendance.verify", args=[out.id]), {"decision": decision, "remarks": remarks})

    def test_sick_employee_never_clocks_in_so_nothing_is_paid(self):
        self.assertEqual(self.ot_pay(), 0.0)
        self.assertIsNone(self.ot.hours)

    def test_leaving_on_time_pays_nothing(self):
        punch_day(self.employee, self.OT_DAY, time_out="17:30")
        self.assertEqual(self.ot_pay(), 0.0)
        self.assertEqual(float(self.ot.hours), 0.0)

    def test_payroll_derives_hours_itself_without_the_overtime_page(self):
        punch_day(self.employee, self.OT_DAY, time_out="19:30")  # 2h past the 5:30 scheduled out
        self.assertEqual(self.ot_pay(), round(2 * self.hourly * 1.25, 2))
        self.assertEqual(float(self.ot.hours), 2.0)

    def test_hours_are_capped_at_the_approved_plan(self):
        punch_day(self.employee, self.OT_DAY, time_out="21:00")  # 3.5h worked, 3h approved
        self.assertEqual(self.ot_pay(), round(3 * self.hourly * 1.25, 2))
        self.assertEqual(float(self.ot.hours), 3.5)

    def test_a_13h_day_is_held_until_hr_verifies_it(self):
        punch_day(self.employee, self.OT_DAY, time_out="22:00")  # 13.5h — looks like a forgotten clock-out
        self.assertEqual(self.ot_pay(), 0.0)
        self.assertIsNone(self.ot.hours)
        self.login(self.tech)
        self.assertContains(self.client.get(reverse("overtime.index")), "awaiting HR check")

    def test_hr_rejecting_the_day_on_the_attendance_log_cancels_the_pay(self):
        punch_day(self.employee, self.OT_DAY, time_out="22:00")
        self.verify("rejected", "Forgot to clock out — left at 6")
        self.ot.refresh_from_db()
        self.assertEqual(float(self.ot.hours), 0.0)  # flipped by the decision itself, before payroll runs
        self.assertEqual(self.ot_pay(), 0.0)
        self.login(self.tech)
        self.assertContains(self.client.get(reverse("overtime.index")), "rejected by HR")

    def test_hr_approving_the_day_releases_the_capped_hours(self):
        punch_day(self.employee, self.OT_DAY, time_out="22:00")
        self.verify("approved")
        self.assertEqual(self.ot_pay(), round(3 * self.hourly * 1.25, 2))
        self.assertEqual(float(self.ot.hours), 4.5)

    def test_changing_the_decision_changes_the_pay(self):
        punch_day(self.employee, self.OT_DAY, time_out="22:00")
        self.verify("approved")
        self.assertGreater(self.ot_pay(), 0)
        self.verify("rejected", "Site lead says they left at 6")
        self.period.refresh_from_db()
        self.assertEqual(self.ot_pay(), 0.0)

    def cancel(self, reason="Out sick that day"):
        self.login(self.hr)
        return self.client.post(reverse("overtime.cancel", args=[self.ot.id]), {"reason": reason})

    def test_hr_cancels_an_approved_ot_the_employee_never_used(self):
        self.assertEqual(self.cancel().status_code, 302)
        self.ot.refresh_from_db()
        self.assertEqual((self.ot.status, self.ot.cancelled_by_id, self.ot.cancel_reason), ("cancelled", self.hr.employee.id, "Out sick that day"))
        self.assertTrue(Notification.objects.filter(user=self.tech, title="Overtime cancelled").exists())
        self.assertEqual(self.ot_pay(), 0.0)
        self.login(self.tech)
        r = self.client.get(reverse("overtime.index"))
        self.assertContains(r, "Cancelled by HR Officer")
        self.assertContains(r, "Out sick that day")

    def test_cancelling_takes_it_off_an_already_computed_line(self):
        punch_day(self.employee, self.OT_DAY, time_out="19:30")
        self.assertGreater(self.ot_pay(), 0)
        self.cancel("Work was moved to next week")
        self.assertEqual(float(self.period.items.get(employee=self.employee).overtime_pay), 0.0)  # no re-run needed

    def test_cannot_cancel_once_the_payslip_is_released(self):
        self.period.released_at = timezone.now()
        self.period.save()
        r = self.cancel()
        self.ot.refresh_from_db()
        self.assertEqual(self.ot.status, "approved")
        self.assertFalse(self.ot.can_cancel())

    def test_only_approvers_see_or_use_cancel(self):
        self.login(self.tech)
        self.assertNotContains(self.client.get(reverse("overtime.index")), "Cancel overtime")
        self.assertEqual(self.client.post(reverse("overtime.cancel", args=[self.ot.id])).status_code, 403)
        self.login(self.hr)
        self.assertContains(self.client.get(reverse("overtime.index")), "Cancel overtime")

    def test_cannot_cancel_once_the_period_is_closed_even_if_unreleased(self):
        self.ot_pay()
        self.login(self.admin)
        self.client.post(reverse("payroll.close", args=[self.period.id]))
        self.ot.refresh_from_db()
        self.assertFalse(self.ot.can_cancel())
        self.cancel()
        self.ot.refresh_from_db()
        self.assertEqual(self.ot.status, "approved")

    def test_approving_after_payroll_was_computed_refreshes_the_line(self):
        self.ot.status = "pending"
        self.ot.save()
        punch_day(self.employee, self.OT_DAY, time_out="19:30")
        self.assertEqual(self.ot_pay(), 0.0)  # computed while still pending
        self.login(self.hr)
        self.client.post(reverse("overtime.approve", args=[self.ot.id]), {"remarks": "ok"})
        self.assertEqual(float(self.period.items.get(employee=self.employee).overtime_pay), round(2 * self.hourly * 1.25, 2))  # no re-run needed

    def test_verifying_a_13h_day_after_payroll_was_computed_refreshes_the_line(self):
        punch_day(self.employee, self.OT_DAY, time_out="22:00")
        self.assertEqual(self.ot_pay(), 0.0)  # held
        self.verify("approved")
        self.assertEqual(float(self.period.items.get(employee=self.employee).overtime_pay), round(3 * self.hourly * 1.25, 2))

    def test_payroll_page_warns_about_overtime_that_will_not_be_paid(self):
        punch_day(self.employee, self.OT_DAY, time_out="22:00")  # approved, but 13h+ and unverified
        other = OvertimeRequest.objects.create(employee=self.hr.employee, ot_date=self.OT_DAY, requested_hours=2, ot_type="regular", reason="Pending one", status="pending")
        self.login(self.admin)
        r = self.client.get(reverse("payroll.index") + f"?period={self.period.id}")
        self.assertContains(r, "will <u>not</u> be paid")
        self.assertContains(r, "1 request still waiting for a decision")
        self.assertContains(r, "1 approved request on a 13h+ day")
        self.verify("approved")
        other.status = "denied"
        other.save()
        self.login(self.admin)
        self.assertNotContains(self.client.get(reverse("payroll.index") + f"?period={self.period.id}"), "will <u>not</u> be paid")

    def test_a_cancelled_date_can_be_filed_again(self):
        d = timezone.now().date() + timedelta(days=7)
        future = OvertimeRequest.objects.create(employee=self.employee, ot_date=d, requested_hours=2, ot_type="regular", reason="Planned site work", status="approved")
        data = {"ot_date": d.isoformat(), "planned_start": "18:00", "planned_end": "20:00", "reason": "Rescheduled pull-out"}
        self.login(self.tech)
        self.assertRedirects(self.client.post(reverse("overtime.store"), data), reverse("overtime.create"), fetch_redirect_response=False)  # already approved for that date
        self.login(self.hr)
        self.client.post(reverse("overtime.cancel", args=[future.id]), {"reason": "Client postponed"})
        self.login(self.tech)
        self.assertRedirects(self.client.post(reverse("overtime.store"), data), reverse("overtime.index"), fetch_redirect_response=False)
        self.assertEqual(self.employee.overtime_requests.filter(ot_date=d).count(), 2)

class RatesTest(AppTestCase):
    def test_update_rates(self):
        self.login(self.admin)
        data = {f"rates[{k}]": v for k, v in rates.DEFAULTS.items()}
        data["rates[allowance_percent]"] = "10"
        r = self.client.post(reverse("payroll.rates.update"), data)
        self.assertRedirects(r, reverse("payroll.rates"), fetch_redirect_response=False)
        self.assertEqual(rates.get("allowance_percent"), 10.0)
        period = PayrollPeriod.ensure_for(date(2026, 9, 1))
        PayrollRunner().generate(period, "test")
        line = period.items.get(employee=self.employee)
        self.assertEqual(float(line.allowances), 1250.0)

    def test_out_of_range_rate_is_refused(self):
        self.login(self.admin)
        data = {f"rates[{k}]": v for k, v in rates.DEFAULTS.items()}
        data["rates[ot_regular_percent]"] = "50"
        self.client.post(reverse("payroll.rates.update"), data)
        self.assertEqual(rates.get("ot_regular_percent"), 125.0)


class DeductionTest(AppTestCase):
    def setUp(self):
        self.period = PayrollPeriod.ensure_for(date(2026, 9, 1))
        self.login(self.admin)

    def test_loan_is_collected_across_cutoffs(self):
        r = self.client.post(reverse("payroll.deductions.store"), {"type": "loan", "description": "Cash advance", "total_amount": "3000", "cutoffs": "3", "starts_on": "2026-09-01", "employee_id": self.employee.id})
        self.assertEqual(r.status_code, 302)
        loan = PayrollDeduction.objects.get()
        self.assertEqual((loan.type, float(loan.installment_amount), float(loan.balance)), ("loan", 1000.0, 3000.0))
        PayrollRunner().generate(self.period, "test")
        line = self.period.items.get(employee=self.employee)
        self.assertEqual(float(line.loan_deduction), 1000.0)
        loan.refresh_from_db()
        self.assertEqual(float(loan.balance), 2000.0)
        # a re-run reverses and re-collects instead of double-charging
        PayrollRunner().generate(self.period, "test", force=True)
        loan.refresh_from_db()
        self.assertEqual(float(loan.balance), 2000.0)

    def test_missing_item_is_split_across_the_people_responsible(self):
        other = Employee.objects.get(user=self.dev)
        r = self.client.post(reverse("payroll.deductions.store"), {"type": "missing_item", "description": "Lost drill", "total_amount": "5000", "cutoffs": "2", "starts_on": "2026-09-01", "site_id": self.project.id, "employees": [self.employee.id, other.id]})
        self.assertEqual(r.status_code, 302)
        rows = PayrollDeduction.objects.filter(type="missing_item").order_by("employee_id")
        self.assertEqual(rows.count(), 2)
        self.assertEqual([float(d.total_amount) for d in rows], [2500.0, 2500.0])
        self.assertEqual([float(d.installment_amount) for d in rows], [1250.0, 1250.0])

    def test_validation(self):
        self.client.post(reverse("payroll.deductions.store"), {"type": "loan", "description": "", "total_amount": "0", "cutoffs": "3", "starts_on": "", "employee_id": ""})
        self.assertEqual(PayrollDeduction.objects.count(), 0)
        r = self.client.get(reverse("payroll.deductions"))
        self.assertContains(r, "Describe it")

    def test_cancel_stops_further_collection(self):
        self.client.post(reverse("payroll.deductions.store"), {"type": "loan", "description": "Cash advance", "total_amount": "3000", "cutoffs": "3", "starts_on": "2026-09-01", "employee_id": self.employee.id})
        loan = PayrollDeduction.objects.get()
        r = self.client.post(reverse("payroll.deductions.cancel", args=[loan.id]), {"reason": "Paid in cash"})
        self.assertEqual(r.status_code, 302)
        loan.refresh_from_db()
        self.assertEqual(loan.status, "cancelled")
        PayrollRunner().generate(self.period, "test")
        self.assertEqual(float(self.period.items.get(employee=self.employee).loan_deduction), 0.0)

    def test_salary_history_page(self):
        r = self.client.get(reverse("employees.salary-history", args=[self.employee.id]))
        self.assertEqual(r.status_code, 200)
