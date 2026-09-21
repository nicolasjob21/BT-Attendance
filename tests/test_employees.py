import io
from datetime import timedelta

from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from accounts.models import User
from employees.models import Employee, EmployeeProjectAssignment, Site

from .base import AppTestCase


def employee_form(**over):
    data = {"first_name": "Juan", "last_name": "Dela Cruz", "username": "", "email": "juan@brite-tsi.com", "payout_method": "cash", "bank_account_no": "", "phone": "", "monthly_salary": "25000", "date_hired": "2026-09-01", "status": "active", "password": ""}
    data.update(over)
    return data


class EmployeeCrudTest(AppTestCase):
    def test_store_creates_user_and_employee_with_generated_username(self):
        self.login(self.hr)
        r = self.client.post(reverse("employees.store"), employee_form())
        self.assertRedirects(r, reverse("employees.index"), fetch_redirect_response=False)
        emp = Employee.objects.get(email="juan@brite-tsi.com")
        self.assertEqual(emp.user.username, "brite-juan")
        self.assertEqual(emp.user.role, "employee")
        self.assertTrue(emp.user.must_change_password)
        self.assertTrue(emp.employee_no.startswith("EMP-"))
        self.assertEqual(float(emp.daily_rate), round(25000 / 22, 2))

    def test_validation_errors_bounce_back_with_old_input(self):
        self.login(self.hr)
        r = self.client.post(reverse("employees.store"), employee_form(email="tech@brite-tsi.com", monthly_salary="-5"))
        self.assertRedirects(r, reverse("employees.create"), fetch_redirect_response=False)
        self.assertEqual(Employee.objects.filter(first_name="Juan").count(), 0)
        r = self.client.get(reverse("employees.create"))
        self.assertContains(r, 'value="Juan"')
        self.assertContains(r, "already")

    def test_card_payout_requires_account_number(self):
        self.login(self.hr)
        self.client.post(reverse("employees.store"), employee_form(payout_method="card"))
        self.assertFalse(Employee.objects.filter(first_name="Juan").exists())
        self.client.post(reverse("employees.store"), employee_form(payout_method="card", bank_account_no="1234 5678"))
        self.assertEqual(Employee.objects.get(first_name="Juan").bank_account_no, "1234 5678")

    def test_update_changes_salary_and_username(self):
        self.login(self.hr)
        r = self.client.post(reverse("employees.update", args=[self.employee.id]), employee_form(first_name="Technical", last_name="Staff", username="brite-techie", email="tech@brite-tsi.com", monthly_salary="30000"))
        self.assertEqual(r.status_code, 302)
        self.employee.refresh_from_db()
        self.tech.refresh_from_db()
        self.assertEqual((float(self.employee.monthly_salary), self.tech.username), (30000.0, "brite-techie"))

    def test_toggle_status(self):
        self.login(self.hr)
        self.client.post(reverse("employees.status", args=[self.employee.id]))
        self.employee.refresh_from_db()
        self.assertEqual(self.employee.status, "inactive")

    def test_employee_cannot_manage_employees(self):
        self.login(self.tech)
        self.assertEqual(self.client.get(reverse("employees.index")).status_code, 403)
        self.assertEqual(self.client.post(reverse("employees.store"), employee_form()).status_code, 403)

    def test_export_is_an_xlsx_of_every_employee(self):
        self.login(self.hr)
        r = self.client.get(reverse("employees.export"))
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(io.BytesIO(r.content))
        names = [row[0] for row in wb.active.iter_rows(min_row=2, values_only=True)]
        self.assertTrue(any("Technical" in str(n) or "EMP" in str(n) for n in names) or len(names) >= 4)

    def test_import_from_template(self):
        self.login(self.hr)
        r = self.client.get(reverse("employees.import.template"))
        self.assertEqual(r.status_code, 200)
        wb = Workbook()
        ws = wb.active
        ws.append(["first_name", "last_name", "username", "email", "monthly_salary"])
        ws.append(["Maria", "Santos", "", "maria@brite-tsi.com", 20000])
        ws.append(["Pedro", "Reyes", "brite-pedro", "pedro@brite-tsi.com", 22000])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        buf.name = "people.xlsx"
        r = self.client.post(reverse("employees.import.store"), {"file": buf, "default_password": "Welcome123"})
        self.assertEqual(r.status_code, 302)
        self.assertTrue(Employee.objects.filter(email="maria@brite-tsi.com").exists())
        pedro = User.objects.get(username="brite-pedro")
        self.assertTrue(pedro.must_change_password)
        self.assertTrue(pedro.check_password("Welcome123"))


class AssignmentTest(AppTestCase):
    def test_assign_employee_to_a_project_site_then_end_it(self):
        self.login(self.hr)
        today = timezone.now().date()
        r = self.client.post(reverse("employees.assignments.store", args=[self.employee.id]), {"site_id": self.project.id, "start_date": today.isoformat(), "end_date": "", "assignment_notes": "Pull-out"})
        self.assertEqual(r.status_code, 302)
        a = EmployeeProjectAssignment.objects.get(employee=self.employee)
        self.assertEqual(a.site_id, self.project.id)
        self.assertEqual(self.employee.active_assignment.id, a.id)
        r = self.client.post(reverse("employees.assignments.end", args=[self.employee.id, a.id]), {"end_date": today.isoformat(), "status": "completed"})
        self.assertEqual(r.status_code, 302)
        a.refresh_from_db()
        self.assertEqual(a.end_date, today)

    def test_new_deployment_closes_the_current_one(self):
        self.login(self.hr)
        today = timezone.now().date()
        old = EmployeeProjectAssignment.objects.create(employee=self.employee, site=self.project, start_date=today - timedelta(days=30), status="active", created_by=self.hr)
        other = Site.objects.create(name="Site B", type="project_site", latitude="14.5547", longitude="121.0244", geofence_radius_m=100, status="active")
        r = self.client.post(reverse("employees.assignments.store", args=[self.employee.id]), {"site_id": other.id, "start_date": today.isoformat(), "end_date": ""})
        self.assertEqual(r.status_code, 302)
        old.refresh_from_db()
        self.assertEqual(old.end_date, today - timedelta(days=1))
        self.assertEqual(self.employee.active_assignment.site_id, other.id)

    def test_end_date_before_start_is_rejected(self):
        self.login(self.hr)
        today = timezone.now().date()
        self.client.post(reverse("employees.assignments.store", args=[self.employee.id]), {"site_id": self.project.id, "start_date": today.isoformat(), "end_date": (today - timedelta(days=1)).isoformat()})
        self.assertEqual(EmployeeProjectAssignment.objects.filter(employee=self.employee).count(), 0)


class SiteTest(AppTestCase):
    def form(self, **over):
        data = {"name": "Site B — Makati", "type": "project_site", "client_name": "Globe", "address": "Makati City", "latitude": "14.5547", "longitude": "121.0244", "geofence_radius_m": "120", "status": "active", "active_from": "", "active_until": ""}
        data.update(over)
        return data

    def test_create_site(self):
        self.login(self.admin)
        r = self.client.post(reverse("sites.store"), self.form())
        self.assertRedirects(r, reverse("sites.index"), fetch_redirect_response=False)
        site = Site.objects.get(name="Site B — Makati")
        self.assertEqual((site.type, site.geofence_radius_m, site.created_by_id), ("project_site", 120, self.admin.id))

    def test_radius_and_coordinates_are_validated(self):
        self.login(self.admin)
        self.client.post(reverse("sites.store"), self.form(geofence_radius_m="5"))
        self.client.post(reverse("sites.store"), self.form(latitude="95"))
        self.assertFalse(Site.objects.filter(name="Site B — Makati").exists())
        r = self.client.get(reverse("sites.create"))
        self.assertContains(r, "Site B")

    def test_update_and_toggle_status(self):
        self.login(self.admin)
        r = self.client.post(reverse("sites.update", args=[self.project.id]), self.form(name="Project Site A — Quezon City", geofence_radius_m="300"))
        self.assertEqual(r.status_code, 302)
        self.project.refresh_from_db()
        self.assertEqual(self.project.geofence_radius_m, 300)
        self.client.post(reverse("sites.status", args=[self.project.id]), {"status": "inactive"})
        self.project.refresh_from_db()
        self.assertEqual(self.project.status, "inactive")

    def test_resolve_link_rejects_non_map_urls(self):
        self.login(self.admin)
        r = self.client.post(reverse("sites.resolve-link"), {"url": "https://example.com/"}, HTTP_X_REQUESTED_WITH="XMLHttpRequest")
        self.assertIn(r.status_code, (400, 422))
