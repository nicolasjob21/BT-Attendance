import io
import re
import secrets
import string
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation

import requests
from django.contrib import messages
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Count, Q, Sum
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from accounts import rbac
from accounts.decorators import permission_required
from accounts.forms import USERNAME_PATTERN
from accounts.models import User
from payroll import rates

from core.support import form_values

from .models import Employee, EmployeeProjectAssignment, Schedule, Site

USERNAME_PREFIX = "brite"


def suggest_username(first_name, reserve=()):
    name = re.sub(r"[^a-z0-9]+", "", (first_name or "").strip().lower()) or "user"
    base = f"{USERNAME_PREFIX}-{name}"
    candidate, n = base, 2
    while candidate in reserve or User.all_objects.filter(username=candidate).exists():
        candidate = f"{base}{n}"
        n += 1
    return candidate


def _office_schedule():
    return Schedule.objects.filter(is_flexible=False).order_by("id").first()


def _daily_rate(monthly):
    return round(float(monthly) / max(1, rates.get("working_days_per_month")), 2)


# ---- employees ----

@permission_required("manage employees")
def index(request):
    search, status = request.GET.get("search", "").strip(), request.GET.get("status", "")
    qs = Employee.objects.select_related("schedule", "user").search(search).order_by("last_name", "first_name")
    if status:
        qs = qs.filter(status=status)
    page = Paginator(qs, 25).get_page(request.GET.get("page"))
    today = timezone.now().date()
    on_project = Employee.objects.active().filter(project_assignments__status="active", project_assignments__start_date__lte=today).filter(Q(project_assignments__end_date__isnull=True) | Q(project_assignments__end_date__gte=today)).distinct().count()
    stats = {
        "active": Employee.objects.active().count(), "on_project": on_project,
        "inactive": Employee.objects.exclude(status="active").count(),
        "payroll": float(Employee.objects.active().aggregate(s=Sum("monthly_salary"))["s"] or 0),
    }
    return render(request, "employees/index.html", {
        "employees": page, "search": search, "status": status, "stats": stats,
        "import_errors": request.session.pop("import_errors", None),
    })


def _validate(request, employee=None):
    d = {k: (request.POST.get(k) or "").strip() for k in ("first_name", "last_name", "username", "email", "phone", "payout_method", "bank_account_no", "monthly_salary", "date_hired", "status", "password")}
    errors = {}
    if not d["first_name"]:
        errors["first_name"] = "First name is required."
    if not d["last_name"]:
        errors["last_name"] = "Last name is required."
    d["username"] = User.normalize_username(d["username"])
    user_id = employee.user_id if employee else None
    if d["username"]:
        if not USERNAME_PATTERN.match(d["username"]):
            errors["username"] = "Use lowercase letters, numbers and single - _ . separators, e.g. brite-juan."
        elif User.all_objects.filter(username=d["username"]).exclude(pk=user_id).exists():
            errors["username"] = "That username is already taken."
    if not d["email"] or "@" not in d["email"]:
        errors["email"] = "Enter a valid e-mail."
    elif User.all_objects.filter(email=d["email"]).exclude(pk=user_id).exists() or Employee.objects.filter(email=d["email"]).exclude(pk=employee.pk if employee else None).exists():
        errors["email"] = "That e-mail is already in use."
    if d["payout_method"] not in ("card", "cash", ""):
        errors["payout_method"] = "Pick a payout method."
    d["payout_method"] = d["payout_method"] or "cash"
    if d["payout_method"] == "card" and not d["bank_account_no"]:
        errors["bank_account_no"] = "Enter the card / account number for card payout."
    try:
        d["monthly_salary"] = Decimal(d["monthly_salary"] or "0")
        if d["monthly_salary"] < 0:
            raise InvalidOperation
    except InvalidOperation:
        errors["monthly_salary"] = "Enter a valid salary."
    if d["date_hired"]:
        try:
            date.fromisoformat(d["date_hired"])
        except ValueError:
            errors["date_hired"] = "Enter a valid date."
    if d["status"] not in ("active", "inactive", "on_leave"):
        errors["status"] = "Pick a status."
    if d["password"] and len(d["password"]) < 8:
        errors["password"] = "The temporary password must be at least 8 characters."
    return d, errors


def _apply(employee, d, user_id):
    employee.user_id = user_id
    employee.first_name, employee.last_name, employee.email, employee.phone = d["first_name"], d["last_name"], d["email"], d["phone"] or None
    employee.payout_method = d["payout_method"]
    employee.bank_account_no = d["bank_account_no"] if d["payout_method"] == "card" else None
    employee.employee_type = "admin"
    employee.schedule = _office_schedule()
    employee.monthly_salary = d["monthly_salary"]
    employee.daily_rate = _daily_rate(d["monthly_salary"])
    employee.date_hired = d["date_hired"] or None
    employee.status = d["status"]


EMPLOYEE_FIELDS = ["first_name", "last_name", "username", "email", "payout_method", "bank_account_no", "phone", "monthly_salary", "date_hired", "status"]


def _employee_form_context(request, employee=None):
    old = request.session.pop("form_old", None) or {}
    return {
        "employee": employee, "errors": request.session.pop("form_errors", None), "old": old,
        "v": form_values(old, employee, EMPLOYEE_FIELDS, aliases={"username": "user.username"}),
        "back_url": reverse("employees.index"), "back_label": "Back to Employees",
    }


@permission_required("manage employees")
def create(request):
    return render(request, "employees/create.html", _employee_form_context(request))


@permission_required("manage employees")
@require_POST
def store(request):
    d, errors = _validate(request)
    if errors:
        request.session["form_errors"], request.session["form_old"] = errors, {k: str(v) for k, v in d.items() if k != "password"}
        return redirect("employees.create")
    temp_password = d["password"] or "".join(secrets.choice(string.ascii_letters + string.digits) for _ in range(10))
    username = d["username"] or suggest_username(d["first_name"])
    with transaction.atomic():
        user = User(name=f"{d['first_name']} {d['last_name']}".strip(), username=username, email=d["email"])
        user.set_temporary_password(temp_password)
        user.save()
        user.set_role(rbac.EMPLOYEE)
        employee = Employee()
        _apply(employee, d, user.id)
        employee.save()
        employee.assign_employee_no()
    messages.success(request, f"Employee added. Username: {username} · temporary password: {temp_password}")
    return redirect("employees.index")


@permission_required("manage employees")
def edit(request, pk):
    employee = get_object_or_404(Employee.objects.select_related("user", "schedule"), pk=pk)
    assignments = employee.project_assignments.select_related("site", "created_by").order_by("-start_date", "-id")
    project_sites = Site.objects.active_on().exclude(type="office").order_by("name")
    return render(request, "employees/edit.html", {
        **_employee_form_context(request, employee),
        "assignments": assignments, "project_sites": project_sites, "active_assignment": employee.active_assignment,
        "today": timezone.now().date().isoformat(),
    })


@permission_required("manage employees")
@require_POST
def update(request, pk):
    employee = get_object_or_404(Employee.objects.select_related("user"), pk=pk)
    d, errors = _validate(request, employee)
    if errors:
        request.session["form_errors"], request.session["form_old"] = errors, {k: str(v) for k, v in d.items() if k != "password"}
        return redirect("employees.edit", pk=pk)
    with transaction.atomic():
        _apply(employee, d, employee.user_id)
        employee.save()
        user = employee.user
        if user:
            user.name, user.email = f"{d['first_name']} {d['last_name']}".strip(), d["email"]
            if d["username"]:
                user.username = d["username"]
            if d["password"]:
                user.set_temporary_password(d["password"])
            user.save()
    messages.success(request, "Employee updated.")
    return redirect("employees.index")


@permission_required("manage employees")
@require_POST
def toggle_status(request, pk):
    employee = get_object_or_404(Employee, pk=pk)
    employee.status = "inactive" if employee.status == "active" else "active"
    employee.save(update_fields=["status"])
    messages.success(request, f"{employee.full_name} is now {employee.status}.")
    return redirect(request.META.get("HTTP_REFERER") or "employees.index")


# ---- import / export ----

def _xlsx_response(wb, filename):
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp


def _style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", start_color="FF0E7490")
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = 18
    ws.freeze_panes = "A2"


@permission_required("export employees")
def export(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    headers = ["employee_no", "first_name", "last_name", "username", "email", "phone", "employee_type", "role", "schedule", "project_site", "monthly_salary", "date_hired", "status"]
    ws.append(headers)
    for e in Employee.objects.select_related("schedule", "user").order_by("last_name"):
        site = e.assigned_site()
        ws.append([e.employee_no, e.first_name, e.last_name, e.user.username if e.user else None, e.email, e.phone, e.employee_type, e.user.role if e.user else None, e.schedule.name if e.schedule else None, site.name if site else None, float(e.monthly_salary), e.date_hired, e.status])
    _style_header(ws, len(headers))
    return _xlsx_response(wb, f"employees-{timezone.now().date():%Y-%m-%d}.xlsx")


@permission_required("manage employees")
def import_form(request):
    return render(request, "employees/import.html", {"errors": request.session.pop("form_errors", None), "back_url": reverse("employees.index"), "back_label": "Back to Employees"})


@permission_required("manage employees")
def import_template(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    ws.append(["first_name", "last_name", "username", "email", "monthly_salary"])
    ws.append(["Juan", "Dela Cruz", "brite-juan", "juan@brite-tsi.com", 25000])
    ws.append(["Maria", "Santos", "", "maria@brite-tsi.com", 20000])
    _style_header(ws, 5)
    return _xlsx_response(wb, "employee-import-template.xlsx")


@permission_required("manage employees")
@require_POST
def import_store(request):
    f = request.FILES.get("file")
    default_password = request.POST.get("default_password", "")
    errors = {}
    if not f or not f.name.lower().endswith((".xlsx", ".xls", ".csv", ".txt")):
        errors["file"] = "Upload an .xlsx or .csv file."
    if len(default_password) < 8:
        errors["default_password"] = "The default password must be at least 8 characters."
    if errors:
        request.session["form_errors"] = errors
        return redirect("employees.import")
    if f.name.lower().endswith((".csv", ".txt")):
        import csv

        rows = list(csv.reader(io.StringIO(f.read().decode("utf-8-sig"))))
    else:
        ws = load_workbook(f, read_only=True, data_only=True).active
        rows = [["" if c is None else c for c in r] for r in ws.iter_rows(values_only=True)]
    header = [re.sub(r"[^a-z0-9]+", "_", str(h or "").strip().lower()).strip("_") for h in (rows[0] if rows else [])]
    office = _office_schedule()
    created = skipped = 0
    problems, used = [], []
    for n, row in enumerate(rows[1:], start=2):
        row = [str(c).strip() for c in row]
        if not any(row):
            continue
        r = dict(zip(header, row))
        email, username = r.get("email", "").strip(), User.normalize_username(r.get("username", ""))
        first, last = r.get("first_name", "").strip(), r.get("last_name", "").strip()
        try:
            salary = float(r.get("monthly_salary") or 0)
        except ValueError:
            salary = 0.0
        if not email or not first or "@" not in email:
            skipped += 1
            problems.append(f"Row {n}: missing/invalid name or email.")
            continue
        if User.all_objects.filter(email=email).exists() or Employee.objects.filter(email=email).exists():
            skipped += 1
            problems.append(f"Row {n}: {email} already exists.")
            continue
        if username:
            if not USERNAME_PATTERN.match(username):
                skipped += 1
                problems.append(f"Row {n}: username '{username}' may only use letters, numbers, - _ . (e.g. brite-juan).")
                continue
            if username in used or User.all_objects.filter(username=username).exists():
                skipped += 1
                problems.append(f"Row {n}: username '{username}' already exists.")
                continue
        else:
            username = suggest_username(first, used)
        used.append(username)
        with transaction.atomic():
            user = User(name=f"{first} {last}".strip(), username=username, email=email)
            user.set_temporary_password(default_password)
            user.save()
            user.set_role(rbac.EMPLOYEE)
            emp = Employee.objects.create(user=user, first_name=first, last_name=last, email=email, employee_type="admin", schedule=office, monthly_salary=salary, daily_rate=_daily_rate(salary), status="active")
            emp.assign_employee_no()
        created += 1
    messages.success(request, f"Imported {created} employees" + (f", skipped {skipped}." if skipped else "."))
    request.session["import_errors"] = problems[:10]
    return redirect("employees.index")


# ---- project assignments ----

@permission_required("manage employees")
@require_POST
def assignment_store(request, pk):
    employee = get_object_or_404(Employee, pk=pk)
    site = Site.objects.filter(pk=request.POST.get("site_id") or 0, status="active").exclude(type="office").first()
    try:
        start = date.fromisoformat(request.POST.get("start_date", ""))
    except ValueError:
        start = None
    end = None
    if request.POST.get("end_date"):
        try:
            end = date.fromisoformat(request.POST.get("end_date"))
        except ValueError:
            end = None
    if not site or not start or (end and end < start):
        messages.error(request, "Pick an active project site and a valid start date.")
        return redirect("employees.edit", pk=pk)
    with transaction.atomic():
        for a in employee.project_assignments.active_on(start):
            a.end(start - timedelta(days=1))
        EmployeeProjectAssignment.objects.create(employee=employee, site=site, start_date=start, end_date=end, status="active", assignment_notes=(request.POST.get("assignment_notes") or "").strip() or None, created_by=request.user)
    messages.success(request, f"{employee.full_name} assigned to {site.name} from {start:%b %-d, %Y}.")
    return redirect("employees.edit", pk=pk)


@permission_required("manage employees")
@require_POST
def assignment_end(request, pk, assignment_id):
    employee = get_object_or_404(Employee, pk=pk)
    a = get_object_or_404(EmployeeProjectAssignment.objects.select_related("site"), pk=assignment_id, employee=employee)
    status = request.POST.get("status") if request.POST.get("status") in ("ended", "cancelled") else "ended"
    on = None
    if request.POST.get("end_date"):
        try:
            on = date.fromisoformat(request.POST.get("end_date"))
        except ValueError:
            on = None
    a.end(on, status)
    messages.success(request, f"Assignment to {a.site.name} {a.status}.")
    return redirect("employees.edit", pk=pk)


# ---- sites ----

@permission_required("manage settings", "manage sites")
def sites_index(request):
    status = request.GET.get("status", "")
    today = timezone.now().date()
    active_filter = Q(assignments__status="active", assignments__start_date__lte=today) & (Q(assignments__end_date__isnull=True) | Q(assignments__end_date__gte=today))
    qs = Site.objects.annotate(active_assignments_count=Count("assignments", filter=active_filter, distinct=True), attendance_logs_count=Count("attendance_logs", distinct=True))
    if status:
        qs = qs.filter(status=status)
    sites = sorted(qs, key=lambda s: (0 if s.type == "office" else 1, 0 if s.status == "active" else 1, s.name))
    all_sites = list(Site.objects.annotate(active_assignments_count=Count("assignments", filter=active_filter, distinct=True)))
    stats = {
        "live": sum(1 for s in all_sites if s.status == "active" and s.type != "office"),
        "deployed": sum(s.active_assignments_count for s in all_sites),
        "ending": sum(1 for s in all_sites if s.status == "active" and s.active_until and today <= s.active_until <= today + timedelta(days=14)),
        "finished": sum(1 for s in all_sites if s.status != "active"),
    }
    return render(request, "sites/index.html", {"sites": sites, "status": status, "stats": stats, "today": today, "soon": today + timedelta(days=14)})


COORD_PATTERNS = [
    re.compile(r"!3d(-?\d+(?:\.\d+)?)!4d(-?\d+(?:\.\d+)?)"),
    re.compile(r"[?&](?:q|query|ll|center|destination)=(-?\d+(?:\.\d+)?)\s*,\s*(-?\d+(?:\.\d+)?)"),
    re.compile(r"@(-?\d+(?:\.\d+)?),(-?\d+(?:\.\d+)?)"),
    re.compile(r"^\s*(-?\d+(?:\.\d+)?)\s*,\s*(-?\d+(?:\.\d+)?)\s*$"),
]


def coords_from_url(url):
    from urllib.parse import unquote

    url = unquote(url)
    for rx in COORD_PATTERNS:
        m = rx.search(url)
        if m:
            lat, lng = float(m.group(1)), float(m.group(2))
            if -90 <= lat <= 90 and -180 <= lng <= 180:
                return lat, lng
    return None


@permission_required("manage settings", "manage sites")
def sites_resolve_link(request):
    """Expand a short Google Maps link server-side and return the coordinates in the final URL."""
    url = (request.GET.get("url") or "").strip()
    if not re.match(r"^https?://", url, re.I):
        return JsonResponse({"ok": False, "message": "Not a link."}, status=422)
    final = url
    try:
        r = requests.get(url, allow_redirects=True, timeout=8, headers={"User-Agent": "BT-Attendance/1.0"})
        final = r.url
    except Exception:
        pass
    coords = coords_from_url(final) or coords_from_url(url)
    return JsonResponse({"ok": bool(coords), "url": final, "lat": coords[0] if coords else None, "lng": coords[1] if coords else None})


def _site_form_data(request, site=None):
    d = {k: (request.POST.get(k) or "").strip() for k in ("name", "type", "client_name", "address", "latitude", "longitude", "geofence_radius_m", "status", "active_from", "active_until")}
    errors = {}
    if not d["name"]:
        errors["name"] = "Name is required."
    elif Site.objects.filter(name=d["name"]).exclude(pk=site.pk if site else None).exists():
        errors["name"] = "A location with that name already exists."
    if d["type"] not in Site.TYPES:
        errors["type"] = "Pick a type."
    try:
        d["latitude"], d["longitude"] = Decimal(d["latitude"]), Decimal(d["longitude"])
        if not (-90 <= d["latitude"] <= 90 and -180 <= d["longitude"] <= 180):
            raise InvalidOperation
    except InvalidOperation:
        errors["latitude"] = "Set the location on the map."
    try:
        d["geofence_radius_m"] = int(d["geofence_radius_m"])
        if not 20 <= d["geofence_radius_m"] <= 5000:
            raise ValueError
    except ValueError:
        errors["geofence_radius_m"] = "Radius must be between 20 and 5000 m."
    if d["status"] not in Site.STATUSES:
        errors["status"] = "Pick a status."
    for k in ("active_from", "active_until"):
        if d[k]:
            try:
                d[k] = date.fromisoformat(d[k])
            except ValueError:
                errors[k] = "Enter a valid date."
        else:
            d[k] = None
    if d["active_from"] and d["active_until"] and d["active_until"] < d["active_from"]:
        errors["active_until"] = "Active until must be on or after active from."
    return d, errors


SITE_FIELDS = ["name", "type", "client_name", "address", "latitude", "longitude", "geofence_radius_m", "status", "active_from", "active_until"]


def _site_form_context(request, site=None):
    old = request.session.pop("form_old", None) or {}
    return {
        "site": site, "errors": request.session.pop("form_errors", None), "old": old,
        "v": form_values(old, site, SITE_FIELDS),
        "back_url": reverse("sites.index"), "back_label": "Back to Locations",
    }


@permission_required("manage settings", "manage sites")
def sites_create(request):
    return render(request, "sites/form.html", _site_form_context(request))


@permission_required("manage settings", "manage sites")
@require_POST
def sites_store(request):
    d, errors = _site_form_data(request)
    if errors:
        request.session["form_errors"], request.session["form_old"] = errors, {k: str(v) if v is not None else "" for k, v in d.items()}
        return redirect("sites.create")
    site = Site.objects.create(created_by=request.user, updated_by=request.user, is_headquarters=d["type"] == "office", **d)
    messages.success(request, f"{site.name} added as an attendance location.")
    return redirect("sites.index")


@permission_required("manage settings", "manage sites")
def sites_edit(request, pk):
    today = timezone.now().date()
    site = get_object_or_404(Site, pk=pk)
    site.active_assignments_count = site.assignments.active_on().count()
    site.attendance_logs_count = site.attendance_logs.count()
    return render(request, "sites/form.html", _site_form_context(request, site))


@permission_required("manage settings", "manage sites")
@require_POST
def sites_update(request, pk):
    site = get_object_or_404(Site, pk=pk)
    d, errors = _site_form_data(request, site)
    if errors:
        request.session["form_errors"], request.session["form_old"] = errors, {k: str(v) if v is not None else "" for k, v in d.items()}
        return redirect("sites.edit", pk=pk)
    for k, v in d.items():
        setattr(site, k, v)
    site.is_headquarters, site.updated_by = d["type"] == "office", request.user
    site.save()
    messages.success(request, f"{site.name} updated.")
    return redirect("sites.index")


@permission_required("manage settings", "manage sites")
@require_POST
def sites_status(request, pk):
    site = get_object_or_404(Site, pk=pk)
    status = request.POST.get("status")
    if status not in Site.STATUSES:
        messages.error(request, "Pick a status.")
        return redirect("sites.index")
    site.status, site.updated_by = status, request.user
    site.save(update_fields=["status", "updated_by"])
    ended = 0
    if status != "active":
        for a in site.assignments.active_on():
            a.end()
            ended += 1
    messages.success(request, f"{site.name} is now {Site.STATUSES[status].lower()}." + (f" {ended} active assignment(s) ended." if ended else ""))
    return redirect(request.META.get("HTTP_REFERER") or "sites.index")
